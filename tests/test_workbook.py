from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.workbook import (
    apply_approved_specification,
    create_project_template_workbook,
    create_sample_investment_workbook,
    inspect_workbook,
)


def create_test_workbook() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Specification"

    worksheet["A1"] = "Project Name"
    worksheet["B1"] = None
    worksheet["A2"] = "Record Count"
    worksheet["B2"] = 10000
    worksheet["B3"] = "=B2*2"

    worksheet.merge_cells("A5:B5")
    worksheet["A5"] = "Instructions"

    original = BytesIO()
    workbook.save(original)
    workbook.close()
    return original.getvalue()


def test_inspection_reads_sheet_names_and_values():
    original = create_test_workbook()

    result = inspect_workbook(original)

    assert result["sheet_names"] == ["Specification"]
    assert result["sheet_count"] == 1

    cells = {
        item["cell"]: item["value"]
        for item in result["sheets"][0]["populated_cells"]
    }

    assert cells["A1"] == "Project Name"
    assert cells["B2"] == 10000
    assert cells["B3"] == "=B2*2"


def test_inspection_detects_formulas_and_merged_cells():
    result = inspect_workbook(create_test_workbook())
    sheet = result["sheets"][0]

    assert {
        "cell": "B3",
        "formula": "=B2*2",
    } in sheet["formulas"]

    assert "A5:B5" in sheet["merged_cells"]


def test_inspection_does_not_modify_original_bytes():
    original = create_test_workbook()
    original_before = original

    inspect_workbook(original)

    assert original == original_before

    reopened = load_workbook(BytesIO(original), data_only=False)
    try:
        assert reopened["Specification"]["B3"].value == "=B2*2"
        assert "A5:B5" in [
            str(item)
            for item in reopened["Specification"].merged_cells.ranges
        ]
    finally:
        reopened.close()


def test_empty_workbook_bytes_are_rejected():
    with pytest.raises(ValueError, match="empty"):
        inspect_workbook(b"")


def test_apply_approved_specification_preserves_original_and_writes_copy():
    original = create_test_workbook()
    original_copy = bytes(original)

    approved = [
        {"sheet_name": "Specification", "cell": "B1", "proposed_value": "Investment Profiling"},
    ]

    modified_bytes = apply_approved_specification(original, approved)

    # Verify original bytes remained untouched
    assert original == original_copy

    # Verify modified bytes contain the new value
    reopened = load_workbook(BytesIO(modified_bytes), data_only=False)
    try:
        assert reopened["Specification"]["B1"].value == "Investment Profiling"
        # Formula in B3 is preserved
        assert reopened["Specification"]["B3"].value == "=B2*2"
    finally:
        reopened.close()


def test_apply_approved_specification_skips_formula_overwrite():
    original = create_test_workbook()
    approved = [
        {"sheet_name": "Specification", "cell": "B3", "proposed_value": 99999},  # B3 is a formula
    ]

    modified_bytes = apply_approved_specification(original, approved)
    reopened = load_workbook(BytesIO(modified_bytes), data_only=False)
    try:
        # B3 formula must not have been overwritten
        assert reopened["Specification"]["B3"].value == "=B2*2"
    finally:
        reopened.close()


def test_create_sample_investment_workbook():
    sample_bytes = create_sample_investment_workbook()
    assert len(sample_bytes) > 0

    inspected = inspect_workbook(sample_bytes)
    assert "Project_Overview" in inspected["sheet_names"]
    assert "Data_Dictionary" in inspected["sheet_names"]
    assert "Business_Rules" in inspected["sheet_names"]

    # Check that formulas are recognized
    dict_sheet = next(s for s in inspected["sheets"] if s["name"] == "Data_Dictionary")
    assert len(dict_sheet["formulas"]) > 0


def test_create_project_template_workbook_generic_domains():
    # Test Agriculture template
    agri_bytes = create_project_template_workbook({
        "project_name": "Kenyan Crop Yield Experiment",
        "domain": "Agriculture",
        "country": "Kenya",
        "dataset_objective": "Model crop yield by soil chemistry",
        "record_count": 5000,
        "random_seed": 12,
    })
    inspected_agri = inspect_workbook(agri_bytes)
    assert "Project_Overview" in inspected_agri["sheet_names"]
    assert "Data_Dictionary" in inspected_agri["sheet_names"]

    # Check Healthcare template
    hc_bytes = create_project_template_workbook({
        "project_name": "UK Clinical Vitals Trial",
        "domain": "Healthcare",
        "country": "United Kingdom",
        "dataset_objective": "Track cardiovascular vitals and risk",
        "record_count": 3000,
        "random_seed": 34,
    })
    inspected_hc = inspect_workbook(hc_bytes)
    assert "Business_Rules" in inspected_hc["sheet_names"]


def test_apply_approved_specification_handles_merged_cells():
    """Verify that writing to a cell coordinate that falls inside a merged range does not raise AttributeError: MergedCell is read-only."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Specification"

    # Merge A1:D2
    ws.merge_cells("A1:D2")

    buf = BytesIO()
    wb.save(buf)
    wb.close()
    raw_bytes = buf.getvalue()

    # B1 is a MergedCell inside A1:D2
    approved = [
        {"sheet_name": "Specification", "cell": "B1", "proposed_value": "Merged Title Entry"},
    ]

    result_bytes = apply_approved_specification(raw_bytes, approved)

    reopened = load_workbook(BytesIO(result_bytes), data_only=False)
    try:
        # The top-left cell A1 must have received the value
        assert reopened["Specification"]["A1"].value == "Merged Title Entry"
    finally:
        reopened.close()


def test_apply_approved_specification_handles_lowercase_and_whitespace_coords():
    """Verify coordinate normalization with lowercase and surrounding whitespace."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Specification"
    ws.merge_cells("B4:D4")

    buf = BytesIO()
    wb.save(buf)
    wb.close()
    raw_bytes = buf.getvalue()

    approved = [
        {"sheet_name": "Specification", "cell": " c4 ", "proposed_value": "Normalized Coordinate Value"},
    ]

    result_bytes = apply_approved_specification(raw_bytes, approved)
    reopened = load_workbook(BytesIO(result_bytes), data_only=False)
    try:
        assert reopened["Specification"]["B4"].value == "Normalized Coordinate Value"
    finally:
        reopened.close()
