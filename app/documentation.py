from io import BytesIO
from typing import Any

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas import GenerationConfig, ValidationReport


def generate_dataset_documentation_docx(
    config: GenerationConfig, df: pd.DataFrame, report: ValidationReport
) -> bytes:
    """Generate a plain-language dataset documentation document in DOCX format."""
    doc = Document()

    # Title
    title = doc.add_heading(f"Synthetic Dataset Documentation: {config.project_name}", level=0)
    title.style.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph(
        f"Domain: {config.domain} | Country: {config.country} | Records: {len(df):,} | Columns: {len(df.columns)}"
    )
    doc.add_paragraph(f"Generated Timestamp: {report.timestamp} | Random Seed: {config.random_seed}")

    doc.add_heading("1. Executive Summary & Objective", level=1)
    doc.add_paragraph(
        f"This dataset was generated to support the objective: '{config.dataset_objective}'. "
        "It provides a statistically realistic and strictly synthetic cohort of client records "
        "without containing genuine Personal Identifiable Information (PII) or customer records."
    )

    doc.add_heading("2. Dataset Overview", level=1)
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.style = "Table Grid"
    hdr = meta_table.rows[0].cells
    hdr[0].text = "Attribute"
    hdr[1].text = "Specification"

    attributes = [
        ("Project Name", config.project_name),
        ("Domain", config.domain),
        ("Country / Jurisdiction", config.country),
        ("Dataset Structure", config.data_structure.capitalize()),
        ("Record Count", f"{len(df):,} records"),
        ("Column Count", f"{len(df.columns)} columns"),
        ("Random Seed", str(config.random_seed)),
        ("Target Variable", next((c.name for c in config.columns if c.is_target), "N/A")),
        ("Noise Policy", f"{config.noise_policy.percentage:.0%} bounded continuous perturbation"),
    ]
    for attr, val in attributes:
        row = meta_table.add_row().cells
        row[0].text = attr
        row[1].text = str(val)

    doc.add_heading("3. Data Dictionary", level=1)
    col_table = doc.add_table(rows=1, cols=5)
    col_table.style = "Table Grid"
    chdr = col_table.rows[0].cells
    chdr[0].text = "Column Name"
    chdr[1].text = "Type"
    chdr[2].text = "Distribution / Formula"
    chdr[3].text = "Noise"
    chdr[4].text = "Description"

    for col in config.columns:
        row = col_table.add_row().cells
        row[0].text = col.name
        row[1].text = col.data_type
        row[2].text = col.formula if col.is_derived else col.distribution
        row[3].text = "Yes" if col.noise_eligible else "No"
        row[4].text = col.description

    doc.add_heading("4. Business Rules & Financial Identities", level=1)
    for c in config.constraints:
        p = doc.add_paragraph()
        p.add_run(f"• {c.name} ({c.type}): ").bold = True
        p.add_run(f"{c.description} [Expression: {c.expression}]")

    doc.add_heading("5. Quality Validation Summary", level=1)
    doc.add_paragraph(
        f"Total Executed Checks: {report.total_checks} | Passed: {report.passed} | "
        f"Warnings: {report.warnings} | Failures: {report.failures}"
    )

    doc.add_heading("6. Assumptions and Limitations", level=1)
    for asm in config.assumptions:
        doc.add_paragraph(f"• {asm}")

    doc.add_paragraph(
        "Notice: This dataset is synthetic and intended strictly for application development, "
        "machine learning prototyping, and algorithmic testing. It does not reflect verified individual "
        "financial status or actual bank accounts."
    )

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def generate_validation_report_docx(config: GenerationConfig, report: ValidationReport) -> bytes:
    """Generate a detailed quality validation report in DOCX format."""
    doc = Document()

    title = doc.add_heading(f"Quality Validation Report: {config.project_name}", level=0)
    title.style.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph(f"Timestamp: {report.timestamp} | Iteration: {report.iteration}")
    doc.add_paragraph(
        f"Summary: {report.passed} PASS | {report.warnings} WARNING | {report.failures} FAIL | {report.not_tested} NOT TESTED"
    )

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Check"
    hdr[1].text = "Requirement"
    hdr[2].text = "Actual Result"
    hdr[3].text = "Status"
    hdr[4].text = "Explanation"

    for item in report.items:
        row = table.add_row().cells
        row[0].text = item.check
        row[1].text = item.requirement
        row[2].text = item.actual_result
        row[3].text = item.status
        row[4].text = item.explanation

    if report.unresolved_issues:
        doc.add_heading("Unresolved Issues", level=2)
        for issue in report.unresolved_issues:
            doc.add_paragraph(f"• {issue}")

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def generate_validation_report_pdf(config: GenerationConfig, report: ValidationReport) -> bytes:
    """Generate a clean, professional quality validation report in PDF format."""
    bio = BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    elements = []

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1F497D"),
        spaceAfter=8,
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#444444"),
        spaceAfter=14,
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
    )
    header_cell_style = ParagraphStyle(
        "HeaderCell",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    elements.append(Paragraph(f"Synthetic Data Quality Validation Report", title_style))
    elements.append(
        Paragraph(
            f"<b>Project:</b> {config.project_name} | <b>Timestamp:</b> {report.timestamp} | "
            f"<b>Checks:</b> {report.total_checks} | <b>Passed:</b> {report.passed} | <b>Warnings:</b> {report.warnings} | <b>Failures:</b> {report.failures}",
            meta_style,
        )
    )

    table_data = [[
        Paragraph("Check", header_cell_style),
        Paragraph("Requirement", header_cell_style),
        Paragraph("Actual Result", header_cell_style),
        Paragraph("Status", header_cell_style),
        Paragraph("Explanation", header_cell_style),
    ]]

    for item in report.items:
        # Format status color
        status_color = "#28A745" if item.status == "PASS" else ("#FFC107" if item.status == "WARNING" else "#DC3545")
        status_p = Paragraph(f"<font color='{status_color}'><b>{item.status}</b></font>", cell_style)

        table_data.append([
            Paragraph(item.check, cell_style),
            Paragraph(item.requirement, cell_style),
            Paragraph(item.actual_result, cell_style),
            status_p,
            Paragraph(item.explanation, cell_style),
        ])

    col_widths = [120, 110, 100, 60, 150]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F497D")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F9")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    elements.append(t)

    doc.build(elements)
    return bio.getvalue()


def generate_assumptions_and_limitations_md(config: GenerationConfig, report: ValidationReport) -> str:
    """Generate an explicit Assumptions and Limitations markdown deliverable."""
    lines = [
        f"# Assumptions and Limitations Report",
        f"**Project:** {config.project_name}  ",
        f"**Domain:** {config.domain} | **Jurisdiction:** {config.country}  ",
        f"**Generated:** {report.timestamp} | **Random Seed:** {config.random_seed}  ",
        "",
        "## 1. Domain & Distributional Assumptions",
    ]
    for asm in config.assumptions:
        lines.append(f"- {asm}")

    lines.extend([
        "",
        "## 2. Noise Policy Specification",
        f"- **Configured Target:** {config.noise_policy.percentage:.0%} random perturbation",
        f"- **Strategy:** {config.noise_policy.strategy}",
        f"- **Perturbation Scale:** {config.noise_policy.perturbation_scale:.1%} relative jitter on eligible numerical features",
        f"- **Eligible Features:** {', '.join(config.noise_policy.eligible_columns)}",
        f"- **Strictly Protected Columns:** {', '.join(config.noise_policy.excluded_columns)}",
        "- **Post-Noise Guarantee:** Bounded ranges and derived accounting formulas are strictly re-enforced following perturbation.",
        "",
        "## 3. Data Quality & Validation Summary",
        f"- **Total Checks Executed:** {report.total_checks}",
        f"- **Pass Rate:** {report.passed}/{report.total_checks} ({(report.passed / report.total_checks * 100):.1f}%)" if report.total_checks else "- Checks: None",
        f"- **Warnings:** {report.warnings}",
        f"- **Failures:** {report.failures}",
        "",
        "## 4. Known Limitations & Governance Warnings",
        "- **Synthetic Nature:** This dataset is generated from statistical formulas and copula distributions; it contains no genuine observations or empirical microdata.",
        "- **Privacy & Re-identification:** While no true customer data was copied or sampled, statistical distributions may mirror stylized market archetypes.",
        "- **Regulatory Scope:** This synthetic dataset is intended for software development, model prototyping, and UI testing. It is NOT certified for regulated automated underwriting decisions or statutory capital adequacy computations.",
    ])

    return "\n".join(lines)
